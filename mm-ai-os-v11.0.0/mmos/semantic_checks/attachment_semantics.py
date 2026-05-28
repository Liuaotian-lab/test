from __future__ import annotations
from pathlib import Path
import json, re
from typing import Any
from mmos.kernel.jsonio import write_json, read_json
from mmos.kernel.events import now_iso

try:
    import openpyxl
except Exception:  # pragma: no cover
    openpyxl = None


def _sample_xlsx(path: Path) -> dict[str, Any]:
    if openpyxl is None:
        return {'error':'openpyxl_missing'}
    wb=openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets=[]
    for ws in wb.worksheets:
        rows=[]
        for i,row in enumerate(ws.iter_rows(values_only=True)):
            if i>=8: break
            rows.append([str(x) if x is not None else '' for x in row[:8]])
        sheets.append({'name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'sample':rows})
    return {'sheets':sheets}


def _role_from_name_and_sample(path: Path, sample: dict[str, Any]) -> tuple[str, list[str], list[str]]:
    name=path.name
    text=name+' '+json.dumps(sample, ensure_ascii=False)
    req=[]; risks=[]
    if re.search(r'凸轮|cam|极角|极径|边缘', text, re.I):
        return 'cam_edge_curve', ['cam_plunger_chamber_model','cam_volume_mapping_verifier'], ['angle_degree_vs_radian','cam_contact_point_max_y']
    if re.search(r'针阀|needle|升程|喷油', text, re.I):
        return 'needle_lift_curve', ['injector_needle_valve_model','needle_effective_area_verifier'], ['lift_unit_mm','effective_area_min_of_gap_and_nozzle']
    if re.search(r'弹性模量|elastic|压力|压强|modulus', text, re.I):
        return 'elastic_modulus_pressure_curve', ['density_pressure_relation','density_pressure_integration_verifier'], ['integral_relation_not_linear_constant']
    if re.search(r'result|输出|模板', text, re.I):
        return 'submission_output_template', ['output_schema_mapper'], []
    return 'unknown_tabular_data', [], ['semantic_role_uncertain']


def attachment_semantic_audit(case_dir: Path, strict: bool = False, write: bool = True) -> dict[str, Any]:
    case_dir=Path(case_dir).resolve()
    raw=case_dir/'data'/'raw'
    attachments=[]; failures=[]; warnings=[]
    for p in sorted(raw.glob('*')):
        if p.name.startswith('.') or p.is_dir():
            continue
        if p.suffix.lower() not in {'.xlsx','.xls','.csv','.tsv'}:
            continue
        sample=_sample_xlsx(p) if p.suffix.lower() in {'.xlsx','.xls'} else {'note':'csv/tsv semantic sampling not implemented'}
        role, req, risks=_role_from_name_and_sample(p, sample)
        item={'file':str(p.relative_to(case_dir)),'semantic_role':role,'required_downstream_models':req,'unit_risks':risks,'sample':sample}
        attachments.append(item)
        if role == 'unknown_tabular_data':
            warnings.append({'code':'UNKNOWN_ATTACHMENT_SEMANTIC_ROLE','file':item['file']})
    text=(case_dir/'workspace'/'problem_corpus.md').read_text(encoding='utf-8', errors='ignore') if (case_dir/'workspace'/'problem_corpus.md').exists() else ''
    required_roles=[]
    if '凸轮' in text: required_roles.append('cam_edge_curve')
    if '针阀' in text: required_roles.append('needle_lift_curve')
    if '弹性模量' in text: required_roles.append('elastic_modulus_pressure_curve')
    found={a['semantic_role'] for a in attachments}
    for r in required_roles:
        if r not in found:
            failures.append({'code':'REQUIRED_ATTACHMENT_ROLE_NOT_FOUND','role':r})
    # Check solver diagnostics claim use of roles
    for qdir in sorted((case_dir/'results').glob('Q*')):
        sol=read_json(qdir/'outputs'/'solution_real.json', {}) or {}
        diag=sol.get('diagnostics') if isinstance(sol.get('diagnostics'), dict) else {}
        used=set(diag.get('attachment_roles_used') or [])
        if strict and required_roles and not set(required_roles).issubset(used) and qdir.name in {'Q2','Q3'}:
            failures.append({'code':'SOLVER_DOES_NOT_CLAIM_REQUIRED_ATTACHMENT_ROLES','question_id':qdir.name,'required_roles':required_roles,'used':sorted(used)})
    status='failed' if failures else ('warning' if warnings else 'passed')
    report={'gate':'attachment-semantic-audit','status':status,'strict':strict,'attachments':attachments,'required_roles':required_roles,'failures':failures,'warnings':warnings,'generated_at':now_iso()}
    if write:
        out=case_dir/'workspace'/'attachment_semantics'
        out.mkdir(parents=True, exist_ok=True)
        for a in attachments:
            stem=Path(a['file']).stem
            write_json(out/f'{stem}.semantic.json', a)
        write_json(out/'case_attachment_semantics.json', report)
        write_json(case_dir/'quality'/'attachment_semantic_audit.json', report)
        write_json(case_dir/'.agent'/'gate_reports'/'attachment_semantic_audit.json', report)
    return report
