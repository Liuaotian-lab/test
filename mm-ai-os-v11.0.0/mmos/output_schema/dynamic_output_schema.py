from __future__ import annotations
from pathlib import Path
from typing import Any

from mmos.kernel.jsonio import read_json, write_json
from mmos.kernel.events import now_iso
from mmos.kernel_gates.common import status_from


def _xlsx_info(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
        wb=load_workbook(path, read_only=True, data_only=False)
        sheets=[]
        for ws in wb.worksheets:
            headers=[]
            for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), values_only=True):
                vals=[v for v in row if v is not None]
                if vals:
                    headers=[str(v) for v in vals]
                    break
            sheets.append({'sheet_name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'headers':headers})
        return {'path':str(path),'type':'xlsx','sheets':sheets}
    except Exception as exc:
        return {'path':str(path),'type':'xlsx','error':str(exc)}


def output_template_infer(case_dir: Path) -> dict[str, Any]:
    case_dir=Path(case_dir)
    templates=[]
    for base in [case_dir/'data'/'raw', case_dir]:
        if base.exists():
            for p in sorted(base.rglob('*')):
                if p.is_file() and p.suffix.lower() in {'.xlsx','.xlsm'}:
                    templates.append(_xlsx_info(p))
                elif p.is_file() and p.suffix.lower() in {'.csv'}:
                    templates.append({'path':str(p),'type':'csv','headers':p.read_text(encoding='utf-8',errors='ignore').splitlines()[:1]})
    report={'command':'output-template-infer','status':'passed','templates':templates,'template_count':len(templates),'generated_at':now_iso()}
    write_json(case_dir/'workspace'/'output_schema'/'candidate'/'output_template_inference.json', report)
    return report


def output_schema_synthesize(case_dir: Path) -> dict[str, Any]:
    case_dir=Path(case_dir)
    inf=read_json(case_dir/'workspace'/'output_schema'/'candidate'/'output_template_inference.json', {}) or output_template_infer(case_dir)
    artifacts=[]
    for t in inf.get('templates',[]):
        path=Path(t.get('path',''))
        name=path.name or 'output.xlsx'
        artifacts.append({
            'artifact':name,
            'source_template':t.get('path'),
            'required': True if name.lower().startswith('result') else False,
            'artifact_type':t.get('type'),
            'workbook':{'sheets':t.get('sheets',[])},
            'semantic_checks':['filename_match','sheet_name_match','column_header_match','no_empty_cells','numeric_range_check','freshness_check'] if t.get('type')=='xlsx' else ['shape_check','freshness_check']
        })
    schema={'schema_version':'10.5.dynamic_output_schema','case_id':case_dir.name,'artifacts':artifacts,'generated_at':now_iso()}
    out=case_dir/'workspace'/'output_schema'/'candidate'/'output_schema.json'
    write_json(out,schema)
    return {'status':'passed','path':str(out),'artifact_count':len(artifacts),'schema':schema}


def output_schema_check(case_dir: Path) -> dict[str, Any]:
    case_dir=Path(case_dir)
    schema=read_json(case_dir/'workspace'/'output_schema'/'candidate'/'output_schema.json', {}) or {}
    failures=[]; warnings=[]
    if not schema:
        warnings.append({'code':'MISSING_OUTPUT_SCHEMA','message':'No output schema synthesized; run output-schema-synthesize.'})
    for art in schema.get('artifacts',[]):
        if not art.get('artifact'):
            failures.append({'code':'OUTPUT_SCHEMA_ARTIFACT_NAME_MISSING'})
        if art.get('artifact_type')=='xlsx':
            sheets=(art.get('workbook') or {}).get('sheets') or []
            if not sheets:
                failures.append({'code':'XLSX_SCHEMA_WITHOUT_SHEETS','artifact':art.get('artifact')})
            for s in sheets:
                if not s.get('sheet_name'):
                    failures.append({'code':'SHEET_NAME_MISSING','artifact':art.get('artifact')})
    report={'gate':'output-schema-check','status':status_from(failures,warnings),'failures':failures,'warnings':warnings,'artifact_count':len(schema.get('artifacts',[]) if schema else []),'generated_at':now_iso()}
    write_json(case_dir/'workspace'/'output_schema'/'reports'/'output_schema_check_report.json', report)
    return report


def workbook_schema_validate(case_dir: Path, all_required: bool = False) -> dict[str, Any]:
    case_dir=Path(case_dir)
    schema=read_json(case_dir/'workspace'/'output_schema'/'candidate'/'output_schema.json', {}) or {}
    failures=[]; warnings=[]; checked=[]
    for art in schema.get('artifacts',[]):
        if art.get('artifact_type')!='xlsx':
            continue
        name=art.get('artifact')
        candidates=[case_dir/'final_outputs'/name, case_dir/'package'/name, case_dir/name]
        candidates += list((case_dir/'results').glob(f'*/outputs/{name}')) if (case_dir/'results').exists() else []
        existing=next((p for p in candidates if p.exists()), None)
        if not existing:
            if all_required or art.get('required'):
                failures.append({'code':'REQUIRED_WORKBOOK_MISSING','artifact':name})
            else:
                warnings.append({'code':'OPTIONAL_WORKBOOK_NOT_FOUND','artifact':name})
            continue
        info=_xlsx_info(existing)
        expected={(s.get('sheet_name'), tuple(s.get('headers') or [])) for s in ((art.get('workbook') or {}).get('sheets') or [])}
        actual={(s.get('sheet_name'), tuple(s.get('headers') or [])) for s in info.get('sheets',[]) }
        missing_sheets=[s for s,_ in expected if s and s not in {a for a,_ in actual}]
        if missing_sheets:
            failures.append({'code':'WORKBOOK_MISSING_SHEETS','artifact':name,'missing_sheets':missing_sheets})
        checked.append({'artifact':name,'path':str(existing),'sheet_count':len(info.get('sheets',[])),'status':'failed' if missing_sheets else 'passed'})
    if not schema:
        warnings.append({'code':'NO_OUTPUT_SCHEMA_FOR_WORKBOOK_VALIDATION'})
    report={'gate':'workbook-schema-validate','status':status_from(failures,warnings),'checked':checked,'failures':failures,'warnings':warnings,'generated_at':now_iso()}
    write_json(case_dir/'workspace'/'output_schema'/'reports'/'workbook_schema_validate_report.json', report)
    return report
