from __future__ import annotations
from pathlib import Path
from zipfile import ZipFile
import tarfile
import csv
import re
import xml.etree.ElementTree as ET
from typing import Any
from mmos.kernel.hashing import sha256_file
from mmos.kernel.jsonio import write_json

TEXT_EXT = {'.txt', '.md', '.py', '.json', '.csv', '.tsv', '.yaml', '.yml'}
SPREADSHEET_EXT = {'.xlsx', '.xls', '.csv', '.tsv'}
DOC_EXT = {'.pdf', '.docx', '.txt', '.md'}

TEXT_ENCODINGS = ['utf-8-sig', 'utf-8', 'gb18030', 'gbk', 'big5']

def read_text_smart(path: Path, limit: int | None = None) -> tuple[str, str, str | None]:
    raw = path.read_bytes()
    if limit is not None:
        raw = raw[:limit]
    last_err = None
    for enc in TEXT_ENCODINGS:
        try:
            return raw.decode(enc), enc, None
        except UnicodeDecodeError as e:
            last_err = str(e)
    return raw.decode('utf-8', errors='replace'), 'utf-8-replace', last_err

def normalize_text_table(path: Path, case_dir: Path) -> dict[str, str] | None:
    if path.suffix.lower() not in {'.csv', '.tsv'}:
        return None
    text, enc, err = read_text_smart(path)
    norm_dir = case_dir / 'data' / 'normalized'
    norm_dir.mkdir(parents=True, exist_ok=True)
    target = norm_dir / (path.stem + '.utf8' + path.suffix.lower())
    target.write_text(text, encoding='utf-8')
    out = {'encoding': enc, 'normalized_path': str(target.relative_to(case_dir))}
    if err and enc == 'utf-8-replace':
        out['decode_warning'] = err
    return out


def detect_file_type(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == '.pdf': return 'pdf'
    if ext == '.docx': return 'docx'
    if ext in {'.xlsx', '.xls'}: return 'spreadsheet'
    if ext in {'.csv', '.tsv'}: return 'table'
    if ext in {'.txt', '.md'}: return 'text'
    if ext == '.zip': return 'archive'
    if ext in {'.png', '.jpg', '.jpeg', '.webp'}: return 'image'
    return 'unknown'




def _is_safe_member_path(name: str) -> bool:
    member = Path(name)
    if member.is_absolute():
        return False
    return not any(part == '..' for part in member.parts)


def _safe_extract_zip(path: Path, target: Path) -> list[dict[str, Any]]:
    extracted: list[dict[str, Any]] = []
    with ZipFile(path) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            if not _is_safe_member_path(info.filename):
                extracted.append({'member': info.filename, 'status': 'blocked', 'reason': 'unsafe_member_path'})
                continue
            out = target / info.filename
            out.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, out.open('wb') as dst:
                dst.write(src.read())
            extracted.append({'member': info.filename, 'status': 'extracted', 'path': str(out)})
    return extracted


def _safe_extract_tar(path: Path, target: Path) -> list[dict[str, Any]]:
    extracted: list[dict[str, Any]] = []
    with tarfile.open(path) as tf:
        for member in tf.getmembers():
            if not member.isfile():
                continue
            if not _is_safe_member_path(member.name):
                extracted.append({'member': member.name, 'status': 'blocked', 'reason': 'unsafe_member_path'})
                continue
            out = target / member.name
            out.parent.mkdir(parents=True, exist_ok=True)
            src = tf.extractfile(member)
            if src is None:
                extracted.append({'member': member.name, 'status': 'skipped', 'reason': 'extractfile_none'})
                continue
            with src, out.open('wb') as dst:
                dst.write(src.read())
            extracted.append({'member': member.name, 'status': 'extracted', 'path': str(out)})
    return extracted


def unpack_archives(case_dir: Path) -> dict[str, Any]:
    """Safely unpack archives under data/raw into data/raw/_unpacked.

    Path traversal and absolute archive members are blocked. The function is
    deterministic and writes an audit report before normal corpus ingestion.
    """
    case_dir = Path(case_dir).resolve()
    raw = case_dir / 'data' / 'raw'
    unpack_root = raw / '_unpacked'
    manifest_dir = case_dir / 'data' / 'manifest'
    manifest_dir.mkdir(parents=True, exist_ok=True)
    archives = []
    if raw.exists():
        for path in sorted(raw.rglob('*')):
            if path.is_dir() or '_unpacked' in path.relative_to(raw).parts:
                continue
            suffixes = ''.join(path.suffixes).lower()
            if path.suffix.lower() == '.zip' or suffixes.endswith(('.tar.gz', '.tgz', '.tar.bz2', '.tbz2', '.tar')):
                archives.append(path)
    records: list[dict[str, Any]] = []
    failures: list[dict[str, Any]] = []
    for archive in archives:
        target = unpack_root / archive.stem
        target.mkdir(parents=True, exist_ok=True)
        try:
            if archive.suffix.lower() == '.zip':
                members = _safe_extract_zip(archive, target)
            else:
                members = _safe_extract_tar(archive, target)
            records.append({'archive': str(archive.relative_to(case_dir)), 'target': str(target.relative_to(case_dir)), 'members': members})
        except Exception as exc:
            failures.append({'archive': str(archive.relative_to(case_dir)), 'error': str(exc)})
    report = {'status': 'failed' if failures else 'passed', 'archive_count': len(archives), 'archives': records, 'failures': failures}
    write_json(manifest_dir / 'archive_unpack_report.json', report)
    return report

def extract_docx_text(path: Path) -> str:
    try:
        with ZipFile(path) as z:
            data = z.read('word/document.xml')
        root = ET.fromstring(data)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        texts = [node.text or '' for node in root.findall('.//w:t', ns)]
        return ''.join(texts)
    except Exception as e:
        return f'[DOCX_EXTRACT_FAILED: {e}]'


def extract_pdf_text(path: Path) -> str:
    try:
        import fitz  # type: ignore
        parts = []
        with fitz.open(path) as doc:
            for page in doc:
                parts.append(page.get_text())
        return '\n'.join(parts)
    except Exception:
        pass
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(str(path))
        return '\n'.join(page.extract_text() or '' for page in reader.pages)
    except Exception as e:
        return f'[PDF_EXTRACT_FAILED: {e}]'


def audit_spreadsheet(path: Path) -> dict[str, Any]:
    ext = path.suffix.lower()
    out: dict[str, Any] = {'type': 'spreadsheet'}
    if ext in {'.csv', '.tsv'}:
        delim = '\t' if ext == '.tsv' else ','
        rows = 0; cols = 0; non_empty = 0
        try:
            sample, encoding, decode_warning = read_text_smart(path)
            import io
            out['encoding'] = encoding
            if decode_warning:
                out['decode_warning'] = decode_warning
            with io.StringIO(sample, newline='') as f:
                for row in csv.reader(f, delimiter=delim):
                    rows += 1
                    cols = max(cols, len(row))
                    non_empty += sum(1 for x in row if str(x).strip())
            out.update({'rows': rows, 'cols': cols, 'non_empty_cells': non_empty})
        except Exception as e:
            out['warning'] = str(e)
        return out
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        sheets = []
        for ws in wb.worksheets:
            non_empty = 0
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value not in (None, ''):
                        non_empty += 1
            sheets.append({'name': ws.title, 'rows': ws.max_row, 'cols': ws.max_column, 'non_empty_cells': non_empty})
        out['sheets'] = sheets
    except Exception as e:
        out['warning'] = str(e)
    return out


def file_record(path: Path, base: Path) -> dict[str, Any]:
    rel = str(path.relative_to(base))
    kind = detect_file_type(path)
    rec: dict[str, Any] = {'path': rel, 'type': kind, 'size_bytes': path.stat().st_size, 'sha256': sha256_file(path)}
    if kind in {'text', 'table'} and path.suffix.lower() not in {'.xlsx', '.xls'}:
        try:
            text, enc, err = read_text_smart(path)
            rec.update({'chars': len(text), 'lines': text.count('\n') + 1, 'encoding': enc})
            norm = normalize_text_table(path, base)
            if norm:
                rec.update(norm)
            if err:
                rec['warning'] = err
        except Exception as e:
            rec['warning'] = str(e)
    if kind == 'pdf':
        txt = extract_pdf_text(path)
        rec.update({'extracted_chars': len(txt), 'extract_warning': txt[:120] if txt.startswith('[PDF_EXTRACT_FAILED') else None})
    if kind == 'docx':
        txt = extract_docx_text(path)
        rec.update({'extracted_chars': len(txt), 'extract_warning': txt[:120] if txt.startswith('[DOCX_EXTRACT_FAILED') else None})
    if kind in {'spreadsheet', 'table'}:
        rec.update(audit_spreadsheet(path))
    return rec


def build_problem_corpus(case_dir: Path) -> dict[str, Any]:
    case_dir = Path(case_dir).resolve()
    raw = case_dir / 'data' / 'raw'
    workspace = case_dir / 'workspace'
    manifest_dir = case_dir / 'data' / 'manifest'
    workspace.mkdir(parents=True, exist_ok=True)
    manifest_dir.mkdir(parents=True, exist_ok=True)
    files = []
    corpus_parts = []
    for path in sorted(raw.rglob('*')) if raw.exists() else []:
        if path.is_dir() or path.name == '.gitkeep':
            continue
        rec = file_record(path, case_dir)
        files.append(rec)
        kind = rec['type']
        text = ''
        if kind == 'text':
            text = read_text_smart(path)[0]
        elif kind == 'table' and path.suffix.lower() in {'.csv', '.tsv'}:
            text = read_text_smart(path, limit=20000)[0]
        elif kind == 'docx':
            text = extract_docx_text(path)
        elif kind == 'pdf':
            text = extract_pdf_text(path)
        if text:
            corpus_parts.append(f'\n\n## SOURCE: {rec["path"]}\n\n{text}\n')
    corpus_md = ''.join(corpus_parts).strip() + '\n'
    (workspace / 'problem_corpus.md').write_text(corpus_md, encoding='utf-8')
    data_manifest = {'case_id': case_dir.name, 'files': files}
    write_json(manifest_dir / 'data_manifest.json', data_manifest)
    write_json(workspace / 'problem_corpus.json', {'case_id': case_dir.name, 'sources': files, 'corpus_path': 'workspace/problem_corpus.md'})
    (manifest_dir / 'data_audit.md').write_text(render_audit_md(data_manifest), encoding='utf-8')
    return data_manifest


def render_audit_md(manifest: dict[str, Any]) -> str:
    lines = [f'# Data Audit: {manifest.get("case_id")}\n', '| path | type | size | notes |', '|---|---:|---:|---|']
    for f in manifest.get('files', []):
        notes = []
        if 'sheets' in f: notes.append(f"sheets={len(f['sheets'])}")
        if 'rows' in f: notes.append(f"rows={f['rows']} cols={f.get('cols')}")
        if 'extracted_chars' in f: notes.append(f"chars={f['extracted_chars']}")
        if f.get('warning'): notes.append(f"warning={f['warning']}")
        lines.append(f"| {f['path']} | {f['type']} | {f['size_bytes']} | {'; '.join(notes)} |")
    return '\n'.join(lines) + '\n'
