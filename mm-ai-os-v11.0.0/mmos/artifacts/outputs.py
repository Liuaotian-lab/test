from __future__ import annotations
from pathlib import Path
import json, csv
from mmos.kernel.jsonio import read_json, write_json
from mmos.kernel.hashing import sha256_file
from mmos.kernel.paths import resolve_case_path

TYPE_BY_EXT = {
    '.json': 'json', '.csv': 'csv', '.tsv': 'csv', '.xlsx': 'excel', '.xls': 'excel',
    '.md': 'markdown', '.txt': 'text', '.png': 'figure', '.jpg': 'figure', '.jpeg': 'figure', '.zip': 'archive'
}


def discover_outputs(case_dir: Path, write_draft: bool = False) -> dict:
    case_dir = Path(case_dir).resolve()
    # Discover only deliverable-like outputs, not internal diagnostics or
    # quality-gate artifacts under results/*/tournament, sensitivity, etc.
    # Registering mutable verifier artifacts makes output-build/package non-
    # idempotent because final-gate-check recomputes those artifacts.
    roots = [case_dir / 'engineering' / 'results', case_dir / 'final_outputs' / 'submitted_files', case_dir / 'reports']
    results_root = case_dir / 'results'
    if results_root.exists():
        roots.extend(p / 'outputs' for p in sorted(results_root.iterdir()) if p.is_dir() and (p / 'outputs').exists())
    outputs = []
    seen = set()
    for root in roots:
        if not root.exists():
            continue
        for p in sorted(root.rglob('*')):
            if p.is_file() and p.suffix.lower() in TYPE_BY_EXT:
                rel_path = p.relative_to(case_dir)
                rel = str(rel_path)
                # output-build copies source-relative files under
                # final_outputs/submitted_files/<source-rel>.  On repeated
                # discover/build cycles, do not rediscover those generated
                # copies when the original source still exists; otherwise the
                # registry grows with mirror entries and output-build reports
                # false target collisions. Manual final-only files remain
                # discoverable.
                submitted_prefix = Path('final_outputs') / 'submitted_files'
                try:
                    source_rel = rel_path.relative_to(submitted_prefix)
                except ValueError:
                    source_rel = None
                if source_rel is not None and (case_dir / source_rel).exists():
                    continue
                if rel in seen:
                    continue
                seen.add(rel)
                kind = TYPE_BY_EXT[p.suffix.lower()]
                qid = infer_question_id(p)
                required = is_under(p, case_dir / 'final_outputs')
                outputs.append({
                    'path': rel,
                    'type': kind,
                    'owner_question': qid,
                    'required': required,
                    'validator': default_validator(kind),
                    'sha256': sha256_file(p),
                    'freshness_sources': infer_freshness_sources(case_dir, p, qid)
                })
    registry = {'case_id': case_dir.name, 'outputs': outputs}
    if write_draft:
        write_json(case_dir / 'registry' / 'outputs_registry.draft.json', registry)
    return registry


def is_under(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except Exception:
        return False


def infer_question_id(path: Path) -> str | None:
    parts = [x.upper() for x in path.parts]
    for part in parts:
        if part.startswith('Q') and part[1:].isdigit():
            return part
    name = path.name.upper()
    import re
    for pat in [r'Q(\d+)', r'RESULT(\d+)', r'问题\s*(\d+)']:
        m = re.search(pat, name)
        if m:
            return f'Q{int(m.group(1))}'
    return None


def infer_freshness_sources(case_dir: Path, p: Path, qid: str | None) -> list[str]:
    if not qid:
        return []
    roots = [case_dir / 'results' / qid, case_dir / 'engineering' / 'questions' / qid, case_dir / 'engineering' / 'common']
    out=[]
    for root in roots:
        if root.exists():
            out.append(str(root.relative_to(case_dir)))
    return out


def default_validator(kind: str) -> dict:
    if kind == 'json': return {'schema_required': False, 'required_non_empty': True, 'required_fields': []}
    if kind == 'excel': return {'min_non_empty_cells': 1, 'required_non_empty': True, 'must_open': True}
    if kind == 'csv': return {'min_rows': 1, 'min_cols': 1, 'required_non_empty': True}
    if kind in {'markdown', 'text'}: return {'min_chars': 20, 'required_non_empty': True}
    if kind == 'figure': return {'min_size_bytes': 100, 'must_open': False}
    return {'required_non_empty': True}


def validate_outputs(case_dir: Path, semantic: bool = True, freshness: bool = True) -> dict:
    case_dir = Path(case_dir).resolve()
    registry = read_json(case_dir / 'registry' / 'outputs_registry.json', default=None) or read_json(case_dir / 'registry' / 'outputs_registry.draft.json', {'outputs': []})
    failures = []
    warnings = []
    outputs = registry.get('outputs', []) if isinstance(registry, dict) else []
    if not outputs:
        failures.append({
            'code': 'EMPTY_OUTPUT_REGISTRY',
            'message': 'registry/outputs_registry.json contains no outputs; final delivery cannot be validated.'
        })
    required_count = sum(1 for x in outputs if x.get('required'))
    if required_count == 0:
        failures.append({
            'code': 'NO_REQUIRED_OUTPUTS_DECLARED',
            'message': 'At least one required output must be registered before package-case.'
        })
    seen_paths = set()
    for out in outputs:
        rel_path = out.get('path')
        if not rel_path:
            failures.append({'code': 'OUTPUT_WITHOUT_PATH', 'output': out})
            continue
        if rel_path in seen_paths:
            failures.append({'code': 'DUPLICATE_OUTPUT_PATH', 'path': rel_path})
        seen_paths.add(rel_path)
        try:
            p = resolve_case_path(case_dir, rel_path)
        except ValueError as exc:
            failures.append({'code': 'OUTPUT_PATH_ESCAPES_CASE', 'path': rel_path, 'message': str(exc)})
            continue
        validator = out.get('validator', {}) or {}
        if out.get('required') and not p.exists():
            failures.append({'code': 'MISSING_REQUIRED_OUTPUT', 'path': rel_path})
            continue
        if not p.exists():
            warnings.append({'code': 'OPTIONAL_OUTPUT_NOT_FOUND', 'path': rel_path})
            continue
        if validator.get('required_non_empty') and p.stat().st_size == 0:
            failures.append({'code': 'EMPTY_OUTPUT', 'path': rel_path})
            continue
        kind = out.get('type')
        if semantic:
            if kind == 'json':
                failures.extend(validate_json_output(p, out))
            elif kind == 'csv':
                failures.extend(validate_csv_output(p, out))
            elif kind == 'excel':
                failures.extend(validate_excel_output(p, out))
            elif kind in {'markdown','text'}:
                min_chars = validator.get('min_chars', 1)
                if len(p.read_text(encoding='utf-8', errors='ignore')) < min_chars:
                    failures.append({'code': 'TEXT_TOO_SHORT', 'path': rel_path, 'min_chars': min_chars})
            elif kind == 'figure':
                if p.stat().st_size < validator.get('min_size_bytes', 1):
                    failures.append({'code': 'FIGURE_TOO_SMALL', 'path': rel_path})
        if freshness:
            stale = check_freshness(case_dir, p, out.get('freshness_sources', []), out.get('sha256'))
            if stale:
                warnings.append(stale)
    report = {'gate': 'output-validate', 'status': 'failed' if failures else ('warning' if warnings else 'passed'), 'failures': failures, 'warnings': warnings}
    write_json(case_dir / 'final_outputs' / 'output_validation_report.json', report)
    write_json(case_dir / '.agent' / 'gate_reports' / 'output_validate.json', report)
    return report


def validate_json_output(p: Path, out: dict) -> list[dict]:
    failures=[]
    try:
        data=json.loads(p.read_text(encoding='utf-8'))
    except Exception as e:
        return [{'code':'INVALID_JSON','path':out['path'],'message':str(e)}]
    for field in out.get('validator', {}).get('required_fields', []):
        if field not in data:
            failures.append({'code':'JSON_MISSING_FIELD','path':out['path'],'field':field})
    if 'quality_level' in data and data.get('quality_level') in {'global_optimal','certified_optimal'}:
        diag=data.get('diagnostics',{}) or {}
        if data.get('quality_level') == 'global_optimal' and not diag.get('global_optimality_certificate'):
            failures.append({'code':'UNSUPPORTED_GLOBAL_OPTIMAL_OUTPUT','path':out['path']})
    return failures


def validate_csv_output(p: Path, out: dict) -> list[dict]:
    failures=[]
    try:
        with p.open('r', encoding='utf-8-sig', newline='') as f:
            rows=list(csv.reader(f))
    except Exception as e:
        return [{'code':'INVALID_CSV','path':out['path'],'message':str(e)}]
    min_rows=out.get('validator',{}).get('min_rows',1); min_cols=out.get('validator',{}).get('min_cols',1)
    if len(rows) < min_rows: failures.append({'code':'CSV_TOO_FEW_ROWS','path':out['path'],'rows':len(rows),'min_rows':min_rows})
    if rows and max(len(r) for r in rows) < min_cols: failures.append({'code':'CSV_TOO_FEW_COLS','path':out['path'],'min_cols':min_cols})
    return failures


def validate_excel_output(p: Path, out: dict) -> list[dict]:
    failures=[]
    try:
        from openpyxl import load_workbook
        wb=load_workbook(p, read_only=True, data_only=False)
    except Exception as e:
        return [{'code':'INVALID_EXCEL','path':out['path'],'message':str(e)}]
    non_empty=0; sheets=[]
    for ws in wb.worksheets:
        c=0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ''):
                    c += 1
        non_empty += c; sheets.append({'name':ws.title,'non_empty_cells':c,'rows':ws.max_row,'cols':ws.max_column})
    min_non_empty=out.get('validator',{}).get('min_non_empty_cells',1)
    if non_empty < min_non_empty:
        failures.append({'code':'EXCEL_TOO_FEW_NON_EMPTY_CELLS','path':out['path'],'non_empty_cells':non_empty,'min':min_non_empty,'sheets':sheets})
    return failures


def check_freshness(case_dir: Path, p: Path, sources: list[str], old_hash: str | None) -> dict | None:
    # Hash mismatch means the registry draft is stale relative to current output file.
    if old_hash and p.exists() and sha256_file(p) != old_hash:
        return {'code':'OUTPUT_HASH_CHANGED_SINCE_REGISTRY','path':str(p.relative_to(case_dir))}
    return None
