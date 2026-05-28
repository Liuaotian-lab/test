from __future__ import annotations
from pathlib import Path
from mmos.kernel.events import now_iso
from mmos.kernel.jsonio import write_json


def validate_fast_result_xlsx(case_dir: Path, path: str | None = None) -> dict:
    case_dir=Path(case_dir).resolve()
    xlsx = case_dir / (path or 'final_outputs/submitted_files/result.xlsx')
    failures=[]; warnings=[]; sheets=[]
    if not xlsx.exists():
        failures.append({'code':'FAST_RESULT_XLSX_MISSING','path':str(xlsx.relative_to(case_dir))})
    else:
        try:
            from openpyxl import load_workbook
            wb=load_workbook(xlsx, read_only=True, data_only=True)
            for ws in wb.worksheets:
                non_empty=0
                numeric=0
                for row in ws.iter_rows():
                    for c in row:
                        if c.value not in (None,''):
                            non_empty += 1
                            if isinstance(c.value,(int,float)):
                                numeric += 1
                sheets.append({'name':ws.title,'rows':ws.max_row,'cols':ws.max_column,'non_empty_cells':non_empty,'numeric_cells':numeric})
            if not sheets:
                failures.append({'code':'FAST_RESULT_XLSX_NO_SHEETS','path':str(xlsx.relative_to(case_dir))})
            if sum(s['non_empty_cells'] for s in sheets) < 10:
                failures.append({'code':'FAST_RESULT_XLSX_TOO_SPARSE','path':str(xlsx.relative_to(case_dir)),'sheets':sheets})
            if not any('顶点' in s['name'] or '抛物' in s['name'] or s['numeric_cells'] >= 3 for s in sheets):
                warnings.append({'code':'FAST_VERTEX_SHEET_NOT_DETECTED','message':'未明显检测到理想抛物面顶点数据；请检查附件4模板映射。'})
            if not any(s['numeric_cells'] >= 100 for s in sheets):
                warnings.append({'code':'FAST_NODE_OR_ACTUATOR_ROWS_NOT_DETECTED','message':'未明显检测到大量节点/促动器数值行；请检查是否完整写入300m口径内节点和伸缩量。'})
        except Exception as e:
            failures.append({'code':'FAST_RESULT_XLSX_INVALID','path':str(xlsx.relative_to(case_dir)),'message':str(e)})
    report={'gate':'fast-result-validate','status':'failed' if failures else ('warning' if warnings else 'passed'),'path':str(xlsx.relative_to(case_dir)) if xlsx.is_absolute() else str(xlsx),'sheets':sheets,'failures':failures,'warnings':warnings,'generated_at':now_iso()}
    write_json(case_dir/'.agent'/'gate_reports'/'fast_result_validate.json', report)
    write_json(case_dir/'quality'/'fast_result_validate.json', report)
    return report
